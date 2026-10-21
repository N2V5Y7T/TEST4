#!/usr/bin/env python3
import bz2
import csv
import gzip
import hashlib
import io
import json
import lzma
import math
import os
import re
import shutil
import stat
import subprocess
import sys
import tarfile
import tempfile
import time
import uuid
import zipfile
from pathlib import Path

import numpy as np
from PIL import Image, ImageChops, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps
import pypdf
from pypdf.generic import ArrayObject, DecodedStreamObject, NameObject, NumberObject
import docx
from docx.shared import Inches, Pt
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
EVIDENCE_DIR = ROOT / "docs" / "evidence"
EVIDENCE_DIR.mkdir(parents=True, exist_ok=True)


def _find_ffmpeg() -> str:
    for candidate in (
        ROOT / ("ffmpeg.exe" if os.name == "nt" else "ffmpeg"),
        ROOT / "assets" / ("ffmpeg.exe" if os.name == "nt" else "ffmpeg"),
    ):
        if candidate.exists():
            return str(candidate)
    which_ff = shutil.which("ffmpeg")
    if which_ff:
        return which_ff
    return "/usr/bin/ffmpeg"


def has_ffmpeg() -> bool:
    ff = _find_ffmpeg()
    return bool(ff and Path(ff).exists())


# ============================================================================
# 1. CLEAN-ROOM PURE LOGIC (MIRRORING crates/dewdrop-core & PROMPT.md)
# ============================================================================

class SpringSolver:
    def __init__(self, initial: float, response: float = 0.45, damping_ratio: float = 0.80):
        self.position = float(initial)
        self.velocity = 0.0
        self.target = float(initial)
        self.response = max(0.05, float(response))
        self.damping_ratio = min(max(float(damping_ratio), 0.1), 2.0)

    def set_target(self, target: float):
        self.target = float(target)

    def step(self, dt: float) -> float:
        sub_dt = 1.0 / 240.0
        rem = min(max(dt, 0.0), 0.25)
        omega = 2.0 * math.pi / self.response
        k = omega * omega
        c = 2.0 * self.damping_ratio * omega
        while rem > 1e-9:
            h = min(rem, sub_dt)
            disp = self.position - self.target
            acc = -k * disp - c * self.velocity
            self.velocity += acc * h
            self.position += self.velocity * h
            rem -= h
        return self.position

    def is_settled(self) -> bool:
        return abs(self.position - self.target) < 0.0008 and abs(self.velocity) < 0.005


def compute_ring_placement(cursor_x: int, cursor_y: int, work_area: tuple, dpi_scale: float = 1.0) -> tuple:
    left, top, right, bottom = work_area
    scale = max(0.5, dpi_scale)
    radius = round(230.0 * scale)
    margin = round(48.0 * scale)
    safe_dist = radius + margin

    space_left = cursor_x - left
    space_right = right - cursor_x
    space_top = cursor_y - top
    space_bottom = bottom - cursor_y

    dx_dir = 1.0 if space_right >= space_left else -1.0
    dy_dir = 1.0 if space_bottom >= space_top else -1.0

    h_room = float(max(space_right, space_left, 1))
    v_room = float(max(space_bottom, space_top, 1))
    norm = math.hypot(h_room, v_room)
    raw_cx = cursor_x + round(dx_dir * (h_room / norm) * safe_dist)
    raw_cy = cursor_y + round(dy_dir * (v_room / norm) * safe_dist)

    pad = round(8.0 * scale)
    min_x = left + radius + pad
    max_x = max(min_x, right - radius - pad)
    min_y = top + radius + pad
    max_y = max(min_y, bottom - radius - pad)

    cx = min(max(raw_cx, min_x), max_x)
    cy = min(max(raw_cy, min_y), max_y)

    dist_sq = (cx - cursor_x) ** 2 + (cy - cursor_y) ** 2
    min_clear = radius + (margin // 2)
    if dist_sq < min_clear * min_clear:
        candidates = [
            (cursor_x + safe_dist, cursor_y),
            (cursor_x - safe_dist, cursor_y),
            (cursor_x, cursor_y + safe_dist),
            (cursor_x, cursor_y - safe_dist),
            (cursor_x + safe_dist, cursor_y + safe_dist),
            (cursor_x - safe_dist, cursor_y + safe_dist),
            (cursor_x + safe_dist, cursor_y - safe_dist),
            (cursor_x - safe_dist, cursor_y - safe_dist),
        ]
        best = (cx, cy)
        best_d2 = dist_sq
        for cand_x, cand_y in candidates:
            x = min(max(cand_x, min_x), max_x)
            y = min(max(cand_y, min_y), max_y)
            d2 = (x - cursor_x) ** 2 + (y - cursor_y) ** 2
            if d2 > best_d2:
                best_d2 = d2
                best = (x, y)
        return best
    return (cx, cy)


class RotationPager:
    """
    Tracks cumulative 360° clockwise and counter-clockwise cursor rotation around the ring centre
    with deadzone and hysteresis. Always accumulates rotation angle so the 360° liquid-glass
    progress arc animates smoothly on both single-page and multi-page rings.
    """
    def __init__(self, total_items: int, deadzone_radius: float = 38.0):
        self.current_page = 0
        self.total_pages = max(1, (max(1, total_items) + 5) // 6)
        self.accumulated_rad = 0.0
        self.last_angle = None
        self.deadzone_radius = deadzone_radius
        self.completed_turns = 0

    def update_pointer(self, dx: float, dy: float):
        if math.hypot(dx, dy) < self.deadzone_radius:
            self.last_angle = None
            return None
        angle = math.atan2(dx, -dy)
        if self.last_angle is not None:
            delta = angle - self.last_angle
            while delta > math.pi:
                delta -= 2.0 * math.pi
            while delta < -math.pi:
                delta += 2.0 * math.pi
            if abs(delta) > 0.014:
                self.accumulated_rad += delta
                self.last_angle = angle
        else:
            self.last_angle = angle

        threshold = 2.0 * math.pi + 0.18
        if self.accumulated_rad >= threshold:
            self.accumulated_rad -= 2.0 * math.pi
            self.completed_turns += 1
            if self.total_pages > 1:
                self.current_page = (self.current_page + 1) % self.total_pages
                return self.current_page
            return 0
        elif self.accumulated_rad <= -threshold:
            self.accumulated_rad += 2.0 * math.pi
            self.completed_turns -= 1
            if self.total_pages > 1:
                self.current_page = (self.current_page - 1) % self.total_pages
                return self.current_page
            return 0
        return None


def sanitize_stem(raw_stem: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", raw_stem).strip().rstrip(".").lstrip("-")
    base = cleaned if cleaned else "output"
    prefix = base.split(".")[0].upper()
    reserved = {
        "CON", "PRN", "AUX", "NUL",
        *{f"COM{i}" for i in range(1, 10)},
        *{f"LPT{i}" for i in range(1, 10)},
    }
    if prefix in reserved:
        base = "_" + base
    return base[:140]


_LAST_FALLBACK_USED = False


def resolve_clash_free_output(source_file: Path, new_stem: str, new_ext: str) -> Path:
    """
    Resolves `<source_folder>/Dew Drop/<stem>.<ext>` with automatic `(1)`, `(2)` clash resolution.
    If `<source_folder>/Dew Drop` is not writable (e.g. CD-ROM, locked network share, read-only dir),
    automatically falls back to `~/Documents/Dew Drop` per PROMPT.md §4 and sets `_LAST_FALLBACK_USED`.
    """
    global _LAST_FALLBACK_USED
    _LAST_FALLBACK_USED = False
    dew_dir = source_file.parent / "Dew Drop"
    try:
        dew_dir.mkdir(parents=True, exist_ok=True)
        probe = dew_dir / f".dewdrop-write-probe-{uuid.uuid4().hex[:8]}"
        probe.write_bytes(b"")
        probe.unlink(missing_ok=True)
    except Exception:
        _LAST_FALLBACK_USED = True
        dew_dir = Path.home() / "Documents" / "Dew Drop"
        dew_dir.mkdir(parents=True, exist_ok=True)

    safe = sanitize_stem(new_stem)
    ext = new_ext.lstrip(".")
    cand = dew_dir / (f"{safe}.{ext}" if ext else safe)
    if not cand.exists():
        return cand
    for idx in range(1, 10000):
        cand = dew_dir / (f"{safe} ({idx}).{ext}" if ext else f"{safe} ({idx})")
        if not cand.exists():
            return cand
    raise RuntimeError("Output clash counter exhausted")


def was_output_fallback_used() -> bool:
    return _LAST_FALLBACK_USED


# ============================================================================
# 2. TOOL CATALOGUE, COMPATIBILITY FILTERING & USAGE RANKING (PROMPT.md §4, §6)
# ============================================================================

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif", ".gif", ".ico", ".avif"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
AUDIO_EXTS = {".mp3", ".wav", ".flac", ".m4a", ".aac", ".ogg", ".opus", ".wma"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".webm", ".avi", ".m4v", ".wmv", ".flv", ".ts"}
MEDIA_EXTS = AUDIO_EXTS | VIDEO_EXTS
ARCHIVE_EXTS = {".zip", ".tar", ".gz", ".tgz", ".bz2", ".tbz2", ".xz", ".txz"}
TEXT_EXTS = {".txt", ".md", ".log", ".ini", ".cfg", ".yaml", ".yml", ".xml", ".html", ".htm", ".py", ".js", ".ts", ".css"}

FORMAT_OF_SUFFIX = {
    ".png": ("PNG", ".png"),
    ".jpg": ("JPEG", ".jpg"),
    ".jpeg": ("JPEG", ".jpg"),
    ".webp": ("WEBP", ".webp"),
    ".bmp": ("BMP", ".bmp"),
    ".tiff": ("TIFF", ".tiff"),
    ".tif": ("TIFF", ".tif"),
    ".gif": ("GIF", ".gif"),
    ".ico": ("ICO", ".ico"),
    ".avif": ("WEBP", ".webp"),
}

TOOL_SPECS = {
    # Group 1: Convert image (5)
    "img_to_png":   {"group": "Convert image", "label": "PNG",         "has_options": True,  "desc": "Max width / Rotate",     "exts": IMAGE_EXTS, "exclude_exts": {".png"}, "min_files": 1, "ffmpeg": False},
    "img_to_jpeg":  {"group": "Convert image", "label": "JPEG",        "has_options": True,  "desc": "Quality (1-100)",        "exts": IMAGE_EXTS, "exclude_exts": {".jpg", ".jpeg"}, "min_files": 1, "ffmpeg": False},
    "img_to_webp":  {"group": "Convert image", "label": "WEBP",        "has_options": True,  "desc": "Quality (1-100)",        "exts": IMAGE_EXTS, "exclude_exts": {".webp"}, "min_files": 1, "ffmpeg": False},
    "img_to_gif":   {"group": "Convert image", "label": "GIF",         "has_options": True,  "desc": "Max width / Rotate",     "exts": IMAGE_EXTS, "exclude_exts": {".gif"}, "min_files": 1, "ffmpeg": False},
    "img_to_ico":   {"group": "Convert image", "label": "ICO",         "has_options": True,  "desc": "Square pad (16-256px)",  "exts": IMAGE_EXTS, "exclude_exts": {".ico"}, "min_files": 1, "ffmpeg": False},
    # Group 2: Edit image (7)
    "img_compress":    {"group": "Edit image", "label": "Compress",    "has_options": True,  "desc": "Quality / Target KB",    "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "img_resize":      {"group": "Edit image", "label": "Resize",      "has_options": True,  "desc": "Scale % / Fit box",      "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "img_crop":        {"group": "Edit image", "label": "Crop shape",  "has_options": True,  "desc": "Aspect ratio",           "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "img_rotate_flip": {"group": "Edit image", "label": "Rotate/Flip", "has_options": True,  "desc": "Angle & Mirror",         "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "img_effects":     {"group": "Edit image", "label": "Effects",     "has_options": True,  "desc": "Filter preset",          "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "img_watermark":   {"group": "Edit image", "label": "Watermark",   "has_options": True,  "desc": "Stamp text & corner",    "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "img_strip_meta":  {"group": "Edit image", "label": "Clean EXIF",  "has_options": False, "desc": "Instant EXIF strip",     "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    # Group 3: PDF & docs (13)
    "pdf_from_images":    {"group": "PDF & docs", "label": "Images->PDF", "has_options": True,  "desc": "Page size & Combine",  "exts": IMAGE_EXTS - {".ico"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_merge":          {"group": "PDF & docs", "label": "Merge PDF",   "has_options": False, "desc": "Combine in order",     "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 2, "ffmpeg": False},
    "pdf_split":          {"group": "PDF & docs", "label": "Split PDF",   "has_options": False, "desc": "One file per page",    "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_extract_pages":  {"group": "PDF & docs", "label": "Extract pg",  "has_options": True,  "desc": "Page range (1-3,5)",   "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_delete_pages":   {"group": "PDF & docs", "label": "Delete pg",   "has_options": True,  "desc": "Pages to remove",      "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_reverse":        {"group": "PDF & docs", "label": "Reverse",     "has_options": False, "desc": "Last page first",      "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_rotate":         {"group": "PDF & docs", "label": "Rotate PDF",  "has_options": True,  "desc": "Angle & Page range",   "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_compress":       {"group": "PDF & docs", "label": "Compress",    "has_options": True,  "desc": "Compression strength", "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_to_text":        {"group": "PDF & docs", "label": "PDF->Text",   "has_options": False, "desc": "Extract plain text",   "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_extract_images": {"group": "PDF & docs", "label": "PDF->Pics",   "has_options": False, "desc": "Save embedded images", "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_strip_meta":     {"group": "PDF & docs", "label": "Clean PDF",   "has_options": False, "desc": "Remove scripts/meta",  "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "pdf_to_word":        {"group": "PDF & docs", "label": "PDF->Word",   "has_options": False, "desc": "Convert to .docx",     "exts": PDF_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "word_to_pdf":        {"group": "PDF & docs", "label": "Word->PDF",   "has_options": False, "desc": "Convert .docx to PDF", "exts": DOCX_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    # Group 4: Video & audio (12)
    "media_to_mp4":    {"group": "Video & audio", "label": "MP4",        "has_options": False, "desc": "H.264 + AAC",           "exts": VIDEO_EXTS, "exclude_exts": {".mp4"}, "min_files": 1, "ffmpeg": True},
    "media_to_webm":   {"group": "Video & audio", "label": "WEBM",       "has_options": False, "desc": "VP9 + Opus",            "exts": VIDEO_EXTS, "exclude_exts": {".webm"}, "min_files": 1, "ffmpeg": True},
    "media_to_mp3":    {"group": "Video & audio", "label": "MP3",        "has_options": False, "desc": "Extract / convert MP3", "exts": MEDIA_EXTS, "exclude_exts": {".mp3"}, "min_files": 1, "ffmpeg": True},
    "media_to_m4a":    {"group": "Video & audio", "label": "M4A",        "has_options": False, "desc": "AAC audio",             "exts": MEDIA_EXTS, "exclude_exts": {".m4a"}, "min_files": 1, "ffmpeg": True},
    "media_to_wav":    {"group": "Video & audio", "label": "WAV",        "has_options": False, "desc": "Lossless PCM WAV",      "exts": MEDIA_EXTS, "exclude_exts": {".wav"}, "min_files": 1, "ffmpeg": True},
    "media_to_flac":   {"group": "Video & audio", "label": "FLAC",       "has_options": False, "desc": "Lossless FLAC",         "exts": MEDIA_EXTS, "exclude_exts": {".flac"}, "min_files": 1, "ffmpeg": True},
    "video_compress":  {"group": "Video & audio", "label": "Shrink Vid", "has_options": True,  "desc": "Strength & Max height", "exts": VIDEO_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": True},
    "audio_compress":  {"group": "Video & audio", "label": "Shrink Aud", "has_options": True,  "desc": "Bitrate (64-320 kbps)", "exts": AUDIO_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": True},
    "media_trim":      {"group": "Video & audio", "label": "Trim",       "has_options": True,  "desc": "Start & End time",      "exts": MEDIA_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": True},
    "video_to_gif":    {"group": "Video & audio", "label": "Video->GIF", "has_options": True,  "desc": "Width & Duration",      "exts": VIDEO_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": True},
    "video_mute":      {"group": "Video & audio", "label": "Mute Video", "has_options": False, "desc": "Strip audio track",     "exts": VIDEO_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": True},
    "video_frame_png": {"group": "Video & audio", "label": "Frame->PNG", "has_options": True,  "desc": "Capture timestamp",     "exts": VIDEO_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": True},
    # Group 5: Archives (4)
    "archive_zip":     {"group": "Archives", "label": "Make ZIP", "has_options": True,  "desc": "Compression level",     "exts": None, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "archive_targz":   {"group": "Archives", "label": "TAR.GZ",   "has_options": True,  "desc": "Compression level",     "exts": None, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "archive_gzip":    {"group": "Archives", "label": "GZIP",     "has_options": True,  "desc": "Compression level",     "exts": None, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "archive_extract": {"group": "Archives", "label": "Extract",  "has_options": False, "desc": "Unpack to new folder",  "exts": ARCHIVE_EXTS, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    # Group 6: Data & text (9)
    "json_pretty":   {"group": "Data & text", "label": "JSON Pretty", "has_options": False, "desc": "Indent 2 spaces",      "exts": {".json"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "json_minify":   {"group": "Data & text", "label": "JSON Min",    "has_options": False, "desc": "Compact JSON",         "exts": {".json"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "json_to_csv":   {"group": "Data & text", "label": "JSON->CSV",   "has_options": False, "desc": "Objects to table",     "exts": {".json"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "csv_to_json":   {"group": "Data & text", "label": "CSV->JSON",   "has_options": False, "desc": "Rows to JSON list",    "exts": {".csv", ".tsv"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "text_utf8":     {"group": "Data & text", "label": "To UTF-8",    "has_options": False, "desc": "Normalize UTF-8 + LF", "exts": TEXT_EXTS | {".csv", ".tsv", ".json", ".srt", ".vtt"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "text_crlf":     {"group": "Data & text", "label": "CRLF Lines",  "has_options": False, "desc": "Windows line endings", "exts": TEXT_EXTS | {".csv", ".tsv", ".json", ".srt", ".vtt"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "srt_to_vtt":    {"group": "Data & text", "label": "SRT->VTT",    "has_options": False, "desc": "SubRip to WebVTT",     "exts": {".srt"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "vtt_to_srt":    {"group": "Data & text", "label": "VTT->SRT",    "has_options": False, "desc": "WebVTT to SubRip",     "exts": {".vtt"}, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
    "checksums_all": {"group": "Data & text", "label": "Checksums",   "has_options": False, "desc": "MD5/SHA1/256/512",     "exts": None, "exclude_exts": set(), "min_files": 1, "ffmpeg": False},
}

DEFAULT_GROUP_ORDER = [
    "Convert image",
    "Edit image",
    "PDF & docs",
    "Video & audio",
    "Archives",
    "Data & text",
]


def is_tool_compatible(tool_id: str, staged_files: list[Path], ffmpeg_available: bool = True) -> bool:
    spec = TOOL_SPECS.get(tool_id)
    if not spec:
        return False
    if spec["ffmpeg"] and not ffmpeg_available:
        return False
    if not staged_files:
        return True
    if len(staged_files) < spec["min_files"]:
        return False
    exts = {p.suffix.lower() for p in staged_files}
    if any(p.name.lower().endswith((".tar.gz", ".tar.bz2", ".tar.xz")) for p in staged_files):
        exts.add(".tar")

    # Strict category matching: when user selects Images, show ONLY image tools;
    # when user selects ZIP/Archives, show ONLY Archive tools; PDF/DOCX -> PDF tools; Media -> Media tools.
    data_all_exts = TEXT_EXTS | {".csv", ".tsv", ".json", ".srt", ".vtt"}
    if spec["exts"] is None:
        if exts.issubset(ARCHIVE_EXTS):
            if spec["group"] != "Archives":
                return False
        elif (
            exts.issubset(IMAGE_EXTS)
            or exts.issubset(PDF_EXTS | DOCX_EXTS)
            or exts.issubset(MEDIA_EXTS)
            or exts.issubset(data_all_exts)
        ):
            return False
    else:
        if not exts.issubset(spec["exts"]):
            return False
    if spec["exclude_exts"] and exts.issubset(spec["exclude_exts"]):
        return False
    return True


def filter_and_rank_tools(
    staged_files: list[Path] | None = None,
    ffmpeg_available: bool = True,
    usage_counts: dict[str, int] | None = None,
) -> list[tuple[str, list[tuple[str, str, bool, str]]]]:
    """
    Filters all 50 tools to those compatible with `staged_files` and `ffmpeg_available`,
    and sorts both groups and tools within each group by descending usage count (breaking ties
    by canonical order) per PROMPT.md §4 and §6.
    """
    files = staged_files or []
    usage = usage_counts or {}
    grouped: dict[str, list[tuple[int, int, tuple[str, str, bool, str]]]] = {g: [] for g in DEFAULT_GROUP_ORDER}

    for order_idx, (tid, spec) in enumerate(TOOL_SPECS.items()):
        if not is_tool_compatible(tid, files, ffmpeg_available):
            continue
        gname = spec["group"]
        u_count = int(usage.get(tid, 0))
        entry = (tid, spec["label"], spec["has_options"], spec["desc"])
        grouped[gname].append((-u_count, order_idx, entry))

    result_groups = []
    for g_idx, gname in enumerate(DEFAULT_GROUP_ORDER):
        items = grouped[gname]
        if not items:
            continue
        items.sort(key=lambda x: (x[0], x[1]))
        tools_sorted = [it[2] for it in items]
        group_usage = sum(int(usage.get(t[0], 0)) for t in tools_sorted) + int(usage.get(f"group:{gname}", 0))
        result_groups.append((-group_usage, g_idx, gname, tools_sorted))

    result_groups.sort(key=lambda x: (x[0], x[1]))
    return [(g[2], g[3]) for g in result_groups]


# ============================================================================
# 3. CLEAN-ROOM ENGINE FOR ALL 50 TOOLS + SECURITY GUARDS
# ============================================================================

MAX_IMAGE_PIXELS = 40_000_000
MAX_COMBINED_PIXELS = 120_000_000
MAX_PDF_PAGES = 200
MAX_ARCHIVE_FILES = 2_000
MAX_OUTPUT_BYTES = 512 * 1024 * 1024


def load_safe_image(path: Path) -> Image.Image:
    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    im = Image.open(path)
    w, h = im.size
    if w <= 0 or h <= 0 or (w * h) > MAX_IMAGE_PIXELS:
        raise ValueError(f"Image dimensions {w}x{h} exceed 40 MP ceiling")
    if getattr(im, "n_frames", 1) > 1:
        raise ValueError("Animated multi-frame images are not supported by static image tools")
    im.load()
    try:
        im = ImageOps.exif_transpose(im)
    except Exception:
        pass
    return im


def flatten_alpha_to_rgb(im: Image.Image, bg_color=(255, 255, 255)) -> Image.Image:
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        canvas = Image.new("RGB", rgba.size, bg_color)
        canvas.paste(rgba, mask=rgba.split()[3])
        return canvas
    return im.convert("RGB")


def save_image_in_format(im: Image.Image, out: Path, pil_fmt: str, quality: int = 90):
    fmt = pil_fmt.upper()
    q = max(1, min(100, int(quality)))
    if fmt in ("JPEG", "JPG"):
        flatten_alpha_to_rgb(im).save(out, "JPEG", quality=q, optimize=True, progressive=True)
    elif fmt == "PNG":
        if im.mode not in ("RGB", "RGBA", "L", "LA", "P"):
            im = im.convert("RGBA")
        im.save(out, "PNG", optimize=True, compress_level=9)
    elif fmt == "WEBP":
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGBA" if "A" in im.mode else "RGB")
        im.save(out, "WEBP", quality=q, method=6)
    elif fmt == "GIF":
        if im.mode not in ("P", "L"):
            im = im.convert("RGBA").quantize(colors=256)
        im.save(out, "GIF", optimize=True)
    elif fmt == "BMP":
        flatten_alpha_to_rgb(im).save(out, "BMP")
    elif fmt == "TIFF":
        im.save(out, "TIFF", compression="tiff_deflate")
    elif fmt == "ICO":
        rgba = im.convert("RGBA")
        w, h = rgba.size
        side = max(w, h)
        sq = Image.new("RGBA", (side, side), (0, 0, 0, 0))
        sq.paste(rgba, ((side - w) // 2, (side - h) // 2))
        sq.save(out, "ICO", sizes=[(16, 16), (24, 24), (32, 32), (48, 48), (64, 64), (128, 128), (256, 256)])
    else:
        im.save(out, fmt)


def fit_to_size_bytes(im: Image.Image, fmt: str, target_bytes: int) -> bytes:
    work = flatten_alpha_to_rgb(im) if fmt == "JPEG" else (im.convert("RGBA") if "A" in im.mode else im.convert("RGB"))
    for _ in range(6):
        lo, hi = 5, 95
        best = None
        while lo <= hi:
            mid = (lo + hi) // 2
            buf = io.BytesIO()
            if fmt == "JPEG":
                work.save(buf, "JPEG", quality=mid, optimize=True, progressive=True)
            else:
                work.save(buf, "WEBP", quality=mid, method=6)
            data = buf.getvalue()
            if len(data) <= target_bytes:
                best = data
                lo = mid + 1
            else:
                hi = mid - 1
        if best is not None:
            return best
        nw, nh = max(64, int(work.width * 0.8)), max(64, int(work.height * 0.8))
        if (nw, nh) == work.size:
            break
        work = work.resize((nw, nh), Image.Resampling.LANCZOS)
    buf = io.BytesIO()
    work.save(buf, "JPEG" if fmt == "JPEG" else "WEBP", quality=5)
    return buf.getvalue()


def parse_pages(spec: str, total: int) -> list[int]:
    spec_s = (spec or "").strip()
    if not spec_s:
        raise ValueError("Enter a page range such as 1-3, 5, 8-")
    pages: list[int] = []
    seen = set()
    for part in spec_s.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a_s, b_s = part.split("-", 1)
            a = int(a_s) if a_s.strip() else 1
            b = int(b_s) if b_s.strip() else total
            if a < 1 or b < 1 or a > total:
                raise ValueError(f"Page range '{part}' is out of bounds (1-{total})")
            b = min(b, total)
            step = 1 if a <= b else -1
            for p in range(a, b + step, step):
                if p not in seen:
                    seen.add(p)
                    pages.append(p)
        else:
            p = int(part)
            if p < 1 or p > total:
                raise ValueError(f"Page {p} is out of bounds (1-{total})")
            if p not in seen:
                seen.add(p)
                pages.append(p)
    if not pages:
        raise ValueError("No valid pages selected")
    return pages


def parse_time(spec: str) -> float:
    s = str(spec or "").strip()
    if not s:
        return 0.0
    parts = s.split(":")
    if len(parts) > 3:
        raise ValueError(f"Invalid timestamp: {s}")
    total = 0.0
    for p in parts:
        val = float(p)
        if val < 0:
            raise ValueError(f"Negative time component in {s}")
        total = total * 60.0 + val
    return total


def strip_pdf_active_content(writer: pypdf.PdfWriter):
    """
    Sanitizes PDF catalog & pages against embedded scripts, auto-actions, forms, and malicious annotations
    while preserving safe external http(s)/mailto links without /Next chains (PROMPT.md §6 & §10).
    """
    root = writer._root_object
    for dangerous_key in ("/OpenAction", "/AA", "/JavaScript", "/JS", "/AcroForm", "/Metadata", "/PieceInfo"):
        if dangerous_key in root:
            del root[dangerous_key]
    if "/Names" in root:
        try:
            names_obj = root["/Names"].get_object()
            for k in ("/JavaScript", "/JS", "/EmbeddedFiles"):
                if k in names_obj:
                    del names_obj[k]
        except Exception:
            del root["/Names"]

    for page in writer.pages:
        for k in ("/AA", "/Metadata", "/PieceInfo"):
            if k in page:
                del page[k]
        if "/Annots" in page:
            safe_annots = ArrayObject()
            try:
                raw_annots = page["/Annots"]
                if hasattr(raw_annots, "get_object"):
                    raw_annots = raw_annots.get_object()
                for ref in raw_annots:
                    annot = ref.get_object() if hasattr(ref, "get_object") else ref
                    if annot.get("/Subtype") != "/Link" or "/AA" in annot:
                        continue
                    act = annot.get("/A")
                    if act is not None:
                        act_obj = act.get_object() if hasattr(act, "get_object") else act
                        if act_obj.get("/S") == "/URI" and "/Next" not in act_obj:
                            uri = str(act_obj.get("/URI", "")).strip().lower()
                            if uri.startswith(("http://", "https://", "mailto:")):
                                safe_annots.append(ref)
            except Exception:
                pass
            if safe_annots:
                page[NameObject("/Annots")] = safe_annots
            else:
                del page["/Annots"]


def recompress_pdf_images(writer: pypdf.PdfWriter, quality: int = 70, max_side: int = 1800):
    seen = set()
    for page in writer.pages:
        res = page.get("/Resources")
        if not res:
            continue
        res_obj = res.get_object() if hasattr(res, "get_object") else res
        xobj = res_obj.get("/XObject")
        if not xobj:
            continue
        xobj = xobj.get_object() if hasattr(xobj, "get_object") else xobj
        for key in list(xobj.keys()):
            try:
                obj = xobj[key].get_object()
                if id(obj) in seen or obj.get("/Subtype") != "/Image":
                    continue
                seen.add(id(obj))
                if "/SMask" in obj or "/Mask" in obj:
                    continue
                w = int(obj.get("/Width", 0))
                h = int(obj.get("/Height", 0))
                if w < 64 or h < 64 or (w * h) > MAX_IMAGE_PIXELS:
                    continue
                raw_bytes = obj.get_data()
                im = Image.open(io.BytesIO(raw_bytes))
                im.load()
                if im.mode not in ("RGB", "L"):
                    im = im.convert("RGB")
                if max_side and max(im.size) > max_side:
                    im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
                buf = io.BytesIO()
                im.save(buf, "JPEG", quality=quality, optimize=True)
                jpg_bytes = buf.getvalue()
                if len(jpg_bytes) >= len(raw_bytes):
                    continue
                new_stream = DecodedStreamObject()
                new_stream._data = jpg_bytes
                new_stream[NameObject("/Type")] = NameObject("/XObject")
                new_stream[NameObject("/Subtype")] = NameObject("/Image")
                new_stream[NameObject("/Width")] = NumberObject(im.width)
                new_stream[NameObject("/Height")] = NumberObject(im.height)
                new_stream[NameObject("/ColorSpace")] = NameObject("/DeviceGray" if im.mode == "L" else "/DeviceRGB")
                new_stream[NameObject("/BitsPerComponent")] = NumberObject(8)
                new_stream[NameObject("/Filter")] = NameObject("/DCTDecode")
                xobj[key] = writer._add_object(new_stream)
            except Exception:
                continue


def open_safe_pdf(path: Path) -> pypdf.PdfReader:
    reader = pypdf.PdfReader(str(path))
    if reader.is_encrypted:
        raise ValueError(f"{path.name} is password-protected")
    if len(reader.pages) > MAX_PDF_PAGES:
        raise ValueError(f"{path.name} has {len(reader.pages)} pages (maximum is {MAX_PDF_PAGES})")
    return reader


def execute_tool(tool_id: str, inputs: list[Path], options: dict | None = None) -> Path:
    opts = options or {}
    first = inputs[0]
    stem = first.stem

    # --- GROUP 1: CONVERT IMAGE (5) ---
    if tool_id in ("img_to_png", "img_to_jpeg", "img_to_webp", "img_to_gif", "img_to_ico"):
        im = load_safe_image(first)
        rot = int(opts.get("rotation", 0))
        if rot in (90, 180, 270):
            im = im.rotate(-rot, expand=True)
        max_side = int(opts.get("width", 0))
        if max_side > 0 and max(im.size) > max_side:
            im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)

        fmt_map = {
            "img_to_png": ("png", "PNG"),
            "img_to_jpeg": ("jpg", "JPEG"),
            "img_to_webp": ("webp", "WEBP"),
            "img_to_gif": ("gif", "GIF"),
            "img_to_ico": ("ico", "ICO"),
        }
        ext, pil_fmt = fmt_map[tool_id]
        out = resolve_clash_free_output(first, stem, ext)
        save_image_in_format(im, out, pil_fmt, quality=int(opts.get("quality", 90)))
        return out

    # --- GROUP 2: EDIT IMAGE (7) — KEEP ORIGINAL IMAGE FORMAT (PROMPT.md §6.2) ---
    if tool_id == "img_compress":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("JPEG", ".jpg"))
        if pil_fmt == "BMP":
            pil_fmt, ext_dot = "PNG", ".png"
        max_side = int(opts.get("width", 0))
        if max_side > 0 and max(im.size) > max_side:
            im.thumbnail((max_side, max_side), Image.Resampling.LANCZOS)
        qual = max(1, min(100, int(opts.get("quality", 75))))
        target_kb = int(opts.get("target_kb", 0))
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        if target_kb > 0 and pil_fmt in ("JPEG", "WEBP"):
            out.write_bytes(fit_to_size_bytes(im, pil_fmt, target_kb * 1024))
        elif pil_fmt == "PNG" and qual < 95:
            colors = 256 if qual >= 70 else (128 if qual >= 40 else 64)
            q_im = im.convert("RGBA").quantize(colors=colors, method=Image.Quantize.FASTOCTREE)
            q_im.save(out, "PNG", optimize=True, compress_level=9)
        else:
            save_image_in_format(im, out, pil_fmt, quality=qual)
        return out

    if tool_id == "img_resize":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("PNG", ".png"))
        mode = opts.get("resize_mode", "fit")
        w = int(opts.get("width", 320))
        h = int(opts.get("height", 240))
        if mode == "percent":
            pct = max(1, min(400, int(opts.get("percent", 50)))) / 100.0
            nw, nh = max(1, round(im.width * pct)), max(1, round(im.height * pct))
            if nw * nh > MAX_IMAGE_PIXELS:
                raise ValueError("Resized dimensions exceed 40 MP ceiling")
            im = im.resize((nw, nh), Image.Resampling.LANCZOS)
        elif mode == "exact":
            nw, nh = max(1, w), max(1, h if h > 0 else w)
            if nw * nh > MAX_IMAGE_PIXELS:
                raise ValueError("Resized dimensions exceed 40 MP ceiling")
            im = im.resize((nw, nh), Image.Resampling.LANCZOS)
        elif mode == "fill":
            nw, nh = max(1, w), max(1, h if h > 0 else w)
            if nw * nh > MAX_IMAGE_PIXELS:
                raise ValueError("Resized dimensions exceed 40 MP ceiling")
            im = ImageOps.fit(im, (nw, nh), Image.Resampling.LANCZOS)
        else:
            im.thumbnail((max(1, w), max(1, h if h > 0 else w)), Image.Resampling.LANCZOS)
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        save_image_in_format(im, out, pil_fmt, quality=int(opts.get("quality", 90)))
        return out

    if tool_id == "img_crop":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("PNG", ".png"))
        ratio_str = str(opts.get("ratio", "1:1"))
        rw, rh = [int(x) for x in ratio_str.split(":")]
        target_ratio = rw / rh
        cur_ratio = im.width / im.height
        if cur_ratio > target_ratio:
            new_w = max(1, round(im.height * target_ratio))
            x0 = (im.width - new_w) // 2
            im = im.crop((x0, 0, x0 + new_w, im.height))
        else:
            new_h = max(1, round(im.width / target_ratio))
            y0 = (im.height - new_h) // 2
            im = im.crop((0, y0, im.width, y0 + new_h))
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        save_image_in_format(im, out, pil_fmt, quality=int(opts.get("quality", 92)))
        return out

    if tool_id == "img_rotate_flip":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("PNG", ".png"))
        rot = int(opts.get("rotation", 90))
        if rot in (90, 180, 270):
            im = im.rotate(-rot, expand=True)
        flip = opts.get("flip", "none")
        if flip == "h":
            im = ImageOps.mirror(im)
        elif flip == "v":
            im = ImageOps.flip(im)
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        save_image_in_format(im, out, pil_fmt, quality=int(opts.get("quality", 92)))
        return out

    if tool_id == "img_effects":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("PNG", ".png"))
        alpha = im.split()[-1] if im.mode in ("RGBA", "LA") else None
        rgb = im.convert("RGB")
        effect = opts.get("effect", "grayscale")
        if effect == "grayscale":
            rgb = ImageOps.grayscale(rgb).convert("RGB")
        elif effect == "sepia":
            g = ImageOps.grayscale(rgb)
            rgb = ImageOps.colorize(g, "#2b1d0e", "#f5e1c0")
        elif effect == "invert":
            rgb = ImageOps.invert(rgb)
        elif effect == "autocontrast":
            rgb = ImageOps.autocontrast(rgb, cutoff=1)
        elif effect == "sharpen":
            rgb = ImageEnhance.Sharpness(rgb).enhance(2.0)
        elif effect == "blur":
            rgb = rgb.filter(ImageFilter.GaussianBlur(radius=float(opts.get("blur_px", 3.0))))
        elif effect == "bw":
            g = ImageOps.grayscale(rgb)
            rgb = g.point(lambda p: 255 if p >= 128 else 0).convert("RGB")
        if alpha is not None and pil_fmt in ("PNG", "WEBP", "TIFF"):
            rgb = rgb.convert("RGBA")
            rgb.putalpha(alpha)
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        save_image_in_format(rgb, out, pil_fmt, quality=int(opts.get("quality", 92)))
        return out

    if tool_id == "img_watermark":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("PNG", ".png"))
        rgba = im.convert("RGBA")
        txt = str(opts.get("text", "Confidential")).strip()[:80] or "Dewdrop"
        pos = str(opts.get("position", "br"))
        overlay = Image.new("RGBA", rgba.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(overlay)
        fsize = max(14, min(rgba.width, rgba.height) // 14)
        font = ImageFont.load_default()
        for fp in (
            "C:/Windows/Fonts/segoeuib.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ):
            if Path(fp).exists():
                try:
                    font = ImageFont.truetype(fp, size=fsize)
                    break
                except Exception:
                    pass
        bbox = d.textbbox((0, 0), txt, font=font)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        pad = max(12, min(rgba.width, rgba.height) // 25)
        positions = {
            "br": (rgba.width - tw - pad, rgba.height - th - pad),
            "bl": (pad, rgba.height - th - pad),
            "tr": (rgba.width - tw - pad, pad),
            "tl": (pad, pad),
            "c":  ((rgba.width - tw) // 2, (rgba.height - th) // 2),
        }
        tx, ty = positions.get(pos, positions["br"])
        d.text((tx + 2, ty + 2), txt, font=font, fill=(0, 0, 0, 145))
        d.text((tx, ty), txt, font=font, fill=(255, 255, 255, 215))
        merged = Image.alpha_composite(rgba, overlay)
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        save_image_in_format(merged, out, pil_fmt, quality=int(opts.get("quality", 92)))
        return out

    if tool_id == "img_strip_meta":
        im = load_safe_image(first)
        pil_fmt, ext_dot = FORMAT_OF_SUFFIX.get(first.suffix.lower(), ("PNG", ".png"))
        clean = Image.new(im.mode, im.size)
        clean.putdata(list(im.getdata()))
        out = resolve_clash_free_output(first, stem, ext_dot.lstrip("."))
        save_image_in_format(clean, out, pil_fmt, quality=92)
        return out

    # --- GROUP 3: PDF AND DOCUMENTS (13) ---
    if tool_id == "pdf_from_images":
        combine = bool(opts.get("combine", True))
        page_mode = str(opts.get("page", "A4"))
        rot = int(opts.get("rotation", 0))
        max_w = int(opts.get("width", 0))
        pages_rgb = []
        total_px = 0
        for p in inputs:
            im = load_safe_image(p)
            if rot in (90, 180, 270):
                im = im.rotate(-rot, expand=True)
            if max_w > 0 and max(im.size) > max_w:
                im.thumbnail((max_w, max_w), Image.Resampling.LANCZOS)
            rgb = flatten_alpha_to_rgb(im)
            total_px += rgb.width * rgb.height
            if total_px > MAX_COMBINED_PIXELS:
                raise ValueError("Combined images exceed 120 MP memory budget")
            if page_mode == "A4":
                a4_w, a4_h = (1754, 1240) if rgb.width > rgb.height else (1240, 1754)
                sheet = Image.new("RGB", (a4_w, a4_h), (255, 255, 255))
                fit = rgb.copy()
                fit.thumbnail((a4_w - 96, a4_h - 96), Image.Resampling.LANCZOS)
                sheet.paste(fit, ((a4_w - fit.width) // 2, (a4_h - fit.height) // 2))
                pages_rgb.append(sheet)
            else:
                pages_rgb.append(rgb)

        if combine or len(pages_rgb) == 1:
            out = resolve_clash_free_output(first, stem, "pdf")
            pages_rgb[0].save(out, "PDF", resolution=150.0, save_all=True, append_images=pages_rgb[1:])
            return out
        else:
            out_first = None
            for p_in, page_im in zip(inputs, pages_rgb):
                out = resolve_clash_free_output(p_in, p_in.stem, "pdf")
                page_im.save(out, "PDF", resolution=150.0)
                if out_first is None:
                    out_first = out
            return out_first

    if tool_id == "pdf_merge":
        writer = pypdf.PdfWriter()
        total_pages = 0
        for p in inputs:
            reader = open_safe_pdf(p)
            total_pages += len(reader.pages)
            if total_pages > MAX_PDF_PAGES:
                raise ValueError("Combined PDF exceeds 200-page safety limit")
            for page in reader.pages:
                writer.add_page(page)
        strip_pdf_active_content(writer)
        out = resolve_clash_free_output(first, f"{stem}_merged", "pdf")
        with open(out, "wb") as f:
            writer.write(f)
        return out

    if tool_id == "pdf_split":
        reader = open_safe_pdf(first)
        out_first = None
        for idx, page in enumerate(reader.pages, start=1):
            w = pypdf.PdfWriter()
            w.add_page(page)
            strip_pdf_active_content(w)
            out = resolve_clash_free_output(first, f"{stem}_page_{idx:03d}", "pdf")
            with open(out, "wb") as f:
                w.write(f)
            if out_first is None:
                out_first = out
        return out_first

    if tool_id in ("pdf_extract_pages", "pdf_delete_pages"):
        reader = open_safe_pdf(first)
        n = len(reader.pages)
        selected = parse_pages(str(opts.get("pages", "1")), n)
        writer = pypdf.PdfWriter()
        if tool_id == "pdf_extract_pages":
            for pno in selected:
                writer.add_page(reader.pages[pno - 1])
        else:
            drop_set = set(selected)
            for idx in range(1, n + 1):
                if idx not in drop_set:
                    writer.add_page(reader.pages[idx - 1])
        if len(writer.pages) == 0:
            raise ValueError("Resulting PDF has 0 pages")
        strip_pdf_active_content(writer)
        out = resolve_clash_free_output(first, stem, "pdf")
        with open(out, "wb") as f:
            writer.write(f)
        return out

    if tool_id == "pdf_reverse":
        reader = open_safe_pdf(first)
        writer = pypdf.PdfWriter()
        for page in reversed(reader.pages):
            writer.add_page(page)
        strip_pdf_active_content(writer)
        out = resolve_clash_free_output(first, stem, "pdf")
        with open(out, "wb") as f:
            writer.write(f)
        return out

    if tool_id == "pdf_rotate":
        reader = open_safe_pdf(first)
        n = len(reader.pages)
        angle = int(opts.get("rotation", 90)) or 90
        pages_spec = str(opts.get("pages", "")).strip()
        target_set = set(parse_pages(pages_spec, n)) if pages_spec else set(range(1, n + 1))
        writer = pypdf.PdfWriter()
        for idx, page in enumerate(reader.pages, start=1):
            if idx in target_set:
                page.rotate(angle)
            writer.add_page(page)
        strip_pdf_active_content(writer)
        out = resolve_clash_free_output(first, stem, "pdf")
        with open(out, "wb") as f:
            writer.write(f)
        return out

    if tool_id == "pdf_compress":
        reader = open_safe_pdf(first)
        level = int(opts.get("level", 2))
        writer = pypdf.PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        for page in writer.pages:
            try:
                page.compress_content_streams()
            except Exception:
                pass
        if level >= 2:
            q = 70 if level == 2 else 45
            max_side = 1800 if level == 2 else 1200
            recompress_pdf_images(writer, quality=q, max_side=max_side)
        strip_pdf_active_content(writer)
        writer.add_metadata({})
        out = resolve_clash_free_output(first, stem, "pdf")
        with open(out, "wb") as f:
            writer.write(f)
        return out

    if tool_id == "pdf_to_text":
        reader = open_safe_pdf(first)
        chunks = [page.extract_text() or "" for page in reader.pages]
        out = resolve_clash_free_output(first, stem, "txt")
        out.write_text("\n\n".join(chunks), encoding="utf-8")
        return out

    if tool_id == "pdf_extract_images":
        reader = open_safe_pdf(first)
        saved = None
        for p_idx, page in enumerate(reader.pages, start=1):
            for i_idx, img_file in enumerate(page.images, start=1):
                ext = Path(img_file.name).suffix.lstrip(".") or "png"
                out = resolve_clash_free_output(first, f"{stem}_p{p_idx}_{i_idx}", ext)
                out.write_bytes(img_file.data)
                if saved is None:
                    saved = out
        if saved is None:
            out = resolve_clash_free_output(first, f"{stem}_images", "txt")
            out.write_text("No embedded raster images found in PDF.", encoding="utf-8")
            return out
        return saved

    if tool_id == "pdf_strip_meta":
        reader = open_safe_pdf(first)
        writer = pypdf.PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        strip_pdf_active_content(writer)
        writer.add_metadata({})
        out = resolve_clash_free_output(first, stem, "pdf")
        with open(out, "wb") as f:
            writer.write(f)
        return out

    if tool_id == "pdf_to_word":
        reader = open_safe_pdf(first)
        document = docx.Document()
        total_chars = 0
        total_imgs = 0
        for p_idx, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            total_chars += len(text)
            if p_idx > 0 and (text or len(page.images) > 0):
                document.add_page_break()
            for raw_line in text.splitlines():
                line = raw_line.strip()
                if not line:
                    continue
                if line.isupper() and len(line) < 72:
                    document.add_heading(line, level=1)
                elif line.startswith(("- ", "* ", "• ")):
                    document.add_paragraph(line[2:].strip(), style="List Bullet")
                elif "|" in line and line.count("|") >= 2:
                    cells = [c.strip() for c in line.strip("|").split("|")]
                    tbl = document.add_table(rows=1, cols=len(cells))
                    for c_i, val in enumerate(cells):
                        tbl.rows[0].cells[c_i].text = val
                else:
                    document.add_paragraph(line)
            for img_file in getattr(page, "images", []):
                try:
                    im = Image.open(io.BytesIO(img_file.data))
                    im.load()
                    png_buf = io.BytesIO()
                    im.save(png_buf, "PNG")
                    png_buf.seek(0)
                    w_in = min(5.5, max(1.2, im.width / 96.0))
                    document.add_picture(png_buf, width=Inches(w_in))
                    total_imgs += 1
                except Exception:
                    continue
        if total_chars == 0 and total_imgs == 0:
            raise ValueError("No text layer detected in PDF (scanned PDFs without text layers are not supported)")
        out = resolve_clash_free_output(first, stem, "docx")
        document.save(str(out))
        return out

    if tool_id == "word_to_pdf":
        with zipfile.ZipFile(first, "r") as zf:
            names = [n.lower() for n in zf.namelist()]
            if any("vbaproject.bin" in n for n in names):
                raise ValueError("Refused Word document containing executable VBA macros")
        doc_obj = docx.Document(str(first))
        out = resolve_clash_free_output(first, stem, "pdf")
        c = rl_canvas.Canvas(str(out), pagesize=A4)
        width_pt, height_pt = A4
        y = height_pt - 54
        for para in doc_obj.paragraphs:
            txt = para.text.strip()
            if not txt:
                y -= 10
                continue
            style_name = (para.style.name if para.style else "").lower()
            has_bold = any(bool(r.bold) for r in para.runs)
            has_italic = any(bool(r.italic) for r in para.runs)
            if "heading 1" in style_name:
                c.setFont("Helvetica-Bold", 16)
                step = 22
            elif "heading" in style_name:
                c.setFont("Helvetica-Bold", 13)
                step = 18
            elif has_bold and has_italic:
                c.setFont("Helvetica-BoldOblique", 11)
                step = 15
            elif has_bold:
                c.setFont("Helvetica-Bold", 11)
                step = 15
            elif has_italic:
                c.setFont("Helvetica-Oblique", 11)
                step = 15
            else:
                c.setFont("Helvetica", 11)
                step = 15
            if "list bullet" in style_name:
                txt = f"•  {txt}"
            if y < 60:
                c.showPage()
                y = height_pt - 54
                c.setFont("Helvetica", 11)
            c.drawString(54, y, txt[:105])
            y -= step
        for tbl in doc_obj.tables:
            for row in tbl.rows:
                row_txt = "  |  ".join(cell.text.strip() for cell in row.cells)
                if y < 60:
                    c.showPage()
                    y = height_pt - 54
                c.setFont("Helvetica", 10)
                c.drawString(54, y, row_txt[:105])
                y -= 14
        for rel in doc_obj.part.rels.values():
            if "image" in getattr(rel, "reltype", ""):
                try:
                    blob = rel.target_part.blob
                    im = Image.open(io.BytesIO(blob))
                    im.load()
                    max_w = width_pt - 108
                    scale = min(1.0, max_w / max(1, im.width), 240.0 / max(1, im.height))
                    dw, dh = im.width * scale, im.height * scale
                    if y - dh < 54:
                        c.showPage()
                        y = height_pt - 54
                    c.drawImage(ImageReader(io.BytesIO(blob)), 54, y - dh, width=dw, height=dh)
                    y -= dh + 14
                except Exception:
                    continue
        c.save()
        return out

    # --- GROUP 4: VIDEO AND AUDIO (12, FFMPEG LOCK-DOWN) ---
    if tool_id in (
        "media_to_mp4", "media_to_webm", "media_to_mp3", "media_to_m4a",
        "media_to_wav", "media_to_flac", "video_compress", "audio_compress",
        "media_trim", "video_to_gif", "video_mute", "video_frame_png"
    ):
        is_audio_in = first.suffix.lower() in AUDIO_EXTS
        ext_map = {
            "media_to_mp4": "mp4",
            "media_to_webm": "webm",
            "media_to_mp3": "mp3",
            "media_to_m4a": "m4a",
            "media_to_wav": "wav",
            "media_to_flac": "flac",
            "video_compress": "mp4",
            "audio_compress": "mp3",
            "media_trim": "mp3" if is_audio_in else "mp4",
            "video_to_gif": "gif",
            "video_mute": "mp4",
            "video_frame_png": "png",
        }
        out = resolve_clash_free_output(first, stem, ext_map[tool_id])
        ff_bin = _find_ffmpeg()
        cmd = [
            ff_bin,
            "-nostdin",
            "-y",
            "-hide_banner",
            "-loglevel", "error",
            "-protocol_whitelist", "file,pipe",
            "-threads", "2",
        ]
        start_sec = parse_time(opts.get("start", "")) if opts.get("start") else 0.0
        end_sec = parse_time(opts.get("end", "")) if opts.get("end") else 0.0
        if start_sec > 0:
            cmd += ["-ss", f"{start_sec:.3f}"]
        if end_sec > 0:
            if end_sec <= start_sec:
                raise ValueError("End time must be greater than start time")
            cmd += ["-to", f"{end_sec:.3f}"]
        cmd += ["-i", str(first.resolve())]

        if tool_id == "media_to_mp4":
            cmd += ["-c:v", "libx264", "-preset", "medium", "-crf", "22", "-pix_fmt", "yuv420p", "-c:a", "aac", "-movflags", "+faststart"]
        elif tool_id == "video_compress":
            lvl = int(opts.get("level", 2))
            crf = {1: "24", 2: "28", 3: "33"}.get(lvl, "28")
            maxh = int(opts.get("maxh", 0))
            if maxh > 0:
                cmd += ["-vf", f"scale=-2:'min({maxh},ih)'"]
            cmd += ["-c:v", "libx264", "-preset", "medium", "-crf", crf, "-pix_fmt", "yuv420p", "-c:a", "aac", "-b:a", "128k", "-movflags", "+faststart"]
        elif tool_id == "media_trim":
            if is_audio_in:
                cmd += ["-vn", "-c:a", "libmp3lame", "-b:a", "192k"]
            else:
                cmd += ["-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac"]
        elif tool_id == "media_to_webm":
            cmd += ["-c:v", "libvpx-vp9", "-b:v", "200k", "-c:a", "libopus"]
        elif tool_id in ("media_to_mp3", "audio_compress"):
            cmd += ["-vn", "-c:a", "libmp3lame", "-b:a", f"{int(opts.get('bitrate', 128))}k"]
        elif tool_id == "media_to_m4a":
            cmd += ["-vn", "-c:a", "aac", "-b:a", "192k"]
        elif tool_id == "media_to_wav":
            cmd += ["-vn", "-c:a", "pcm_s16le"]
        elif tool_id == "media_to_flac":
            cmd += ["-vn", "-c:a", "flac"]
        elif tool_id == "video_to_gif":
            gw = max(120, min(800, int(opts.get("width", 320))))
            dur = max(1.0, min(30.0, float(opts.get("duration", 6.0))))
            cmd += [
                "-t", f"{dur:.2f}",
                "-vf", f"fps=12,scale={gw}:-1:flags=lanczos,split[s0][s1];[s0]palettegen=max_colors=128[p];[s1][p]paletteuse=dither=bayer",
            ]
        elif tool_id == "video_mute":
            cmd += ["-an", "-c:v", "libx264", "-pix_fmt", "yuv420p"]
        elif tool_id == "video_frame_png":
            cmd += ["-frames:v", "1"]

        cmd += ["--", str(out.resolve())]
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if res.returncode != 0:
            raise RuntimeError(f"FFmpeg rejected input or failed: {res.stderr.strip()}")
        return out

    # --- GROUP 5: ARCHIVES (4) ---
    if tool_id == "archive_zip":
        lvl = {1: 1, 2: 6, 3: 9}.get(int(opts.get("level", 2)), 6)
        combine = bool(opts.get("combine", True))
        if combine or len(inputs) == 1:
            out = resolve_clash_free_output(first, stem, "zip")
            with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=lvl) as zf:
                for p in inputs:
                    zf.write(p, arcname=sanitize_stem(p.stem) + p.suffix)
            return out
        else:
            out_first = None
            for p in inputs:
                out = resolve_clash_free_output(p, p.stem, "zip")
                with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=lvl) as zf:
                    zf.write(p, arcname=sanitize_stem(p.stem) + p.suffix)
                if out_first is None:
                    out_first = out
            return out_first

    if tool_id == "archive_targz":
        lvl = {1: 1, 2: 6, 3: 9}.get(int(opts.get("level", 2)), 6)
        out = resolve_clash_free_output(first, stem, "tar.gz")
        with gzip.open(out, "wb", compresslevel=lvl) as gz_out:
            with tarfile.open(fileobj=gz_out, mode="w") as tf:
                for p in inputs:
                    info = tf.gettarinfo(str(p), arcname=sanitize_stem(p.stem) + p.suffix)
                    info.uid = 0
                    info.gid = 0
                    info.uname = ""
                    info.gname = ""
                    with open(p, "rb") as fobj:
                        tf.addfile(info, fobj)
        return out

    if tool_id == "archive_gzip":
        lvl = {1: 1, 2: 6, 3: 9}.get(int(opts.get("level", 2)), 6)
        out = resolve_clash_free_output(first, first.name, "gz")
        with open(first, "rb") as f_in, gzip.open(out, "wb", compresslevel=lvl) as f_out:
            shutil.copyfileobj(f_in, f_out)
        return out

    if tool_id == "archive_extract":
        out_folder = resolve_clash_free_output(first, stem, "")
        out_folder.mkdir(parents=True, exist_ok=True)
        name_l = first.name.lower()
        total_written = 0
        file_count = 0

        def check_rel_path(rel_name: str) -> Path:
            norm = rel_name.replace("\\", "/")
            if norm.startswith("/") or ":" in norm:
                raise ValueError(f"Refused archive entry with absolute/drive path: {rel_name}")
            parts = [p for p in norm.split("/") if p and p != "."]
            if not parts or any(p == ".." for p in parts) or len(parts) > 20:
                raise ValueError(f"Refused archive entry with path traversal: {rel_name}")
            safe_parts = [sanitize_stem(Path(p).stem) + Path(p).suffix for p in parts]
            return out_folder.joinpath(*safe_parts)

        if name_l.endswith(".zip"):
            with zipfile.ZipFile(first, "r") as zf:
                for info in zf.infolist():
                    if info.flag_bits & 0x1:
                        raise ValueError("Refused password-protected archive entry")
                    mode_hi = (info.external_attr >> 16) & 0xFFFF
                    if stat.S_ISLNK(mode_hi) or stat.S_ISCHR(mode_hi) or stat.S_ISBLK(mode_hi) or stat.S_ISFIFO(mode_hi) or stat.S_ISSOCK(mode_hi):
                        raise ValueError(f"Refused symlink/device entry in zip: {info.filename}")
                    if info.is_dir():
                        continue
                    file_count += 1
                    if file_count > MAX_ARCHIVE_FILES:
                        raise ValueError("Archive exceeds 2,000 file count limit")
                    dest = check_rel_path(info.filename)
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    with zf.open(info, "r") as src, open(dest, "wb") as dst:
                        while True:
                            chunk = src.read(65536)
                            if not chunk:
                                break
                            total_written += len(chunk)
                            if total_written > MAX_OUTPUT_BYTES:
                                raise ValueError("Archive extraction exceeded 512 MiB limit")
                            dst.write(chunk)
        elif name_l.endswith((".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tbz2", ".tar.xz", ".txz")):
            with tarfile.open(first, "r:*") as tf:
                for member in tf.getmembers():
                    if member.isdir():
                        continue
                    if not member.isfile():
                        raise ValueError("Refused symlink/device entry inside tar archive")
                    file_count += 1
                    if file_count > MAX_ARCHIVE_FILES:
                        raise ValueError("Archive exceeds 2,000 file count limit")
                    dest = check_rel_path(member.name)
                    dest.parent.mkdir(parents=True, exist_ok=True)
                    src = tf.extractfile(member)
                    if src:
                        with open(dest, "wb") as dst:
                            while True:
                                chunk = src.read(65536)
                                if not chunk:
                                    break
                                total_written += len(chunk)
                                if total_written > MAX_OUTPUT_BYTES:
                                    raise ValueError("Archive extraction exceeded 512 MiB limit")
                                dst.write(chunk)
        elif name_l.endswith((".gz", ".bz2", ".xz")):
            opener = gzip.open if name_l.endswith(".gz") else (bz2.open if name_l.endswith(".bz2") else lzma.open)
            dest = out_folder / sanitize_stem(first.stem)
            with opener(first, "rb") as src, open(dest, "wb") as dst:
                while True:
                    chunk = src.read(65536)
                    if not chunk:
                        break
                    total_written += len(chunk)
                    if total_written > MAX_OUTPUT_BYTES:
                        raise ValueError("Decompression exceeded 512 MiB limit")
                    dst.write(chunk)
        else:
            raise ValueError(f"Unsupported archive format: {first.name}")
        return out_folder

    # --- GROUP 6: DATA AND TEXT (9) ---
    if tool_id == "json_pretty":
        obj = json.loads(first.read_text(encoding="utf-8-sig"))
        out = resolve_clash_free_output(first, stem, "json")
        out.write_text(json.dumps(obj, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        return out

    if tool_id == "json_minify":
        obj = json.loads(first.read_text(encoding="utf-8-sig"))
        out = resolve_clash_free_output(first, stem, "json")
        out.write_text(json.dumps(obj, separators=(",", ":"), ensure_ascii=False), encoding="utf-8")
        return out

    if tool_id == "json_to_csv":
        data = json.loads(first.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict):
            for v in data.values():
                if isinstance(v, list) and v and isinstance(v[0], dict):
                    data = v
                    break
            else:
                data = [data]
        if not isinstance(data, list) or not data or not all(isinstance(r, dict) for r in data):
            raise ValueError("JSON to CSV expects a non-empty list of objects")
        cols: list[str] = []
        seen_cols = set()
        for r in data:
            for k in r.keys():
                ks = str(k)
                if ks not in seen_cols:
                    seen_cols.add(ks)
                    cols.append(ks)
        out = resolve_clash_free_output(first, stem, "csv")
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            writer = csv.DictWriter(f, fieldnames=cols)
            writer.writeheader()
            for r in data:
                writer.writerow({
                    k: (json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
                    for k, v in r.items()
                })
        return out

    if tool_id == "csv_to_json":
        text = first.read_text(encoding="utf-8-sig", errors="replace")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except Exception:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        data = [dict(r) for r in reader]
        out = resolve_clash_free_output(first, stem, "json")
        out.write_text(json.dumps(data, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
        return out

    if tool_id == "text_utf8":
        raw = first.read_bytes()
        txt = None
        for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
            try:
                txt = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        if txt is None:
            txt = raw.decode("utf-8", errors="replace")
        txt = txt.replace("\r\n", "\n").replace("\r", "\n")
        out_ext = first.suffix.lstrip(".") or "txt"
        out = resolve_clash_free_output(first, stem, out_ext)
        out.write_text(txt, encoding="utf-8")
        return out

    if tool_id == "text_crlf":
        txt = first.read_text(encoding="utf-8-sig", errors="replace")
        lines = txt.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        out_ext = first.suffix.lstrip(".") or "txt"
        out = resolve_clash_free_output(first, stem, out_ext)
        out.write_bytes("\r\n".join(lines).encode("utf-8"))
        return out

    if tool_id == "srt_to_vtt":
        srt = first.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
        body = re.sub(r"(\d{2}:\d{2}:\d{2}),(\d{3})", r"\1.\2", srt)
        out = resolve_clash_free_output(first, stem, "vtt")
        out.write_text("WEBVTT\n\n" + body.lstrip(), encoding="utf-8")
        return out

    if tool_id == "vtt_to_srt":
        vtt = first.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
        body = re.sub(r"^WEBVTT[^\n]*\n+", "", vtt.lstrip())
        body = re.sub(r"(\d{2}:\d{2}:\d{2})\.(\d{3})", r"\1,\2", body)
        blocks = [b.strip() for b in re.split(r"\n\s*\n", body.strip()) if b.strip()]
        numbered = []
        idx = 1
        for b in blocks:
            lines = b.split("\n")
            if "-->" in lines[0]:
                numbered.append(f"{idx}\n" + "\n".join(lines))
                idx += 1
            elif len(lines) >= 2 and "-->" in lines[1]:
                lines[0] = str(idx)
                numbered.append("\n".join(lines))
                idx += 1
        out = resolve_clash_free_output(first, stem, "srt")
        out.write_text(("\n\n".join(numbered) if numbered else body) + "\n", encoding="utf-8")
        return out

    if tool_id == "checksums_all":
        lines = []
        for p in inputs:
            h_md5 = hashlib.md5()
            h_sha1 = hashlib.sha1()
            h_sha256 = hashlib.sha256()
            h_sha512 = hashlib.sha512()
            sz = 0
            with open(p, "rb") as f:
                while True:
                    chunk = f.read(1024 * 1024)
                    if not chunk:
                        break
                    sz += len(chunk)
                    h_md5.update(chunk)
                    h_sha1.update(chunk)
                    h_sha256.update(chunk)
                    h_sha512.update(chunk)
            lines.extend([
                f"File   : {p.name}",
                f"Size   : {sz:,} bytes",
                f"MD5    : {h_md5.hexdigest()}",
                f"SHA1   : {h_sha1.hexdigest()}",
                f"SHA256 : {h_sha256.hexdigest()}",
                f"SHA512 : {h_sha512.hexdigest()}",
                "",
            ])
        out = resolve_clash_free_output(first, f"{stem}_checksums", "txt")
        out.write_text("\n".join(lines), encoding="utf-8")
        return out

    raise ValueError(f"Unknown tool_id: {tool_id}")


# ============================================================================
# 4. SUBPROCESS WORKER ENTRY POINT (`--worker <job.json> <result.json>`)
# ============================================================================

MULTI_FILE_BATCH_TOOLS = {"pdf_merge", "pdf_from_images", "archive_zip", "archive_targz", "checksums_all"}


def _atomic_write_json(path: Path, payload: dict):
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, path)


def worker_main(job_file: str, result_file: str) -> int:
    """
    Runs a conversion batch in an isolated worker process, writing atomic live status JSON
    to `result_file` after each item so the UI shell can poll progress, suspend/resume the
    worker process via NtSuspendProcess/NtResumeProcess, or terminate it on Cancel.
    """
    res_path = Path(result_file)
    try:
        spec = json.loads(Path(job_file).read_text(encoding="utf-8"))
        tool_id = str(spec["tool_id"])
        files = [Path(p) for p in spec["files"]]
        opts = dict(spec.get("options") or {})
    except Exception as exc:
        _atomic_write_json(res_path, {
            "ok": False, "finished": True, "fraction": 1.0,
            "done": 0, "total": 1, "current": "",
            "outputs": [], "errors": [f"Invalid job specification: {exc}"],
            "fallback_used": False,
        })
        return 2

    outputs: list[str] = []
    errors: list[str] = []
    fallback_any = False

    is_single_call = (
        tool_id in MULTI_FILE_BATCH_TOOLS
        and not (tool_id in ("pdf_from_images", "archive_zip") and not opts.get("combine", True))
    )
    total = 1 if is_single_call else max(1, len(files))

    _atomic_write_json(res_path, {
        "ok": True, "finished": False, "fraction": 0.04,
        "done": 0, "total": total, "current": files[0].name if files else "",
        "outputs": [], "errors": [], "fallback_used": False,
    })

    if is_single_call:
        try:
            out_p = execute_tool(tool_id, files, opts)
            outputs.append(str(out_p))
            fallback_any = fallback_any or was_output_fallback_used()
        except Exception as exc:
            errors.append(f"{files[0].name if files else tool_id}: {exc}")
        _atomic_write_json(res_path, {
            "ok": len(errors) == 0, "finished": True, "fraction": 1.0,
            "done": 1, "total": 1, "current": files[0].name if files else "",
            "outputs": outputs, "errors": errors, "fallback_used": fallback_any,
        })
        return 0 if not errors else 1

    for idx, fpath in enumerate(files):
        _atomic_write_json(res_path, {
            "ok": True, "finished": False,
            "fraction": max(0.05, idx / float(total)),
            "done": idx, "total": total, "current": fpath.name, "current_path": str(fpath),
            "outputs": outputs, "errors": errors, "fallback_used": fallback_any,
        })
        try:
            out_p = execute_tool(tool_id, [fpath], opts)
            outputs.append(str(out_p))
            fallback_any = fallback_any or was_output_fallback_used()
        except Exception as exc:
            errors.append(f"{fpath.name}: {exc}")

    _atomic_write_json(res_path, {
        "ok": len(outputs) > 0, "finished": True, "fraction": 1.0,
        "done": total, "total": total, "current": files[-1].name if files else "",
        "outputs": outputs, "errors": errors, "fallback_used": fallback_any,
    })
    return 0 if len(outputs) > 0 else 1


# ============================================================================
# 5. RUN COMPLETE TEST SUITE (M1, M2, M3, M4, M5, M6, M8, CLEAN-ROOM AUDIT)
# ============================================================================

def run_all_tests():
    results = {
        "unit_tests": {},
        "liquid_glass_optics_and_zero_orange_audit": {},
        "tools_50": {},
        "pdf_to_word_10_samples": {},
        "batch_100": {},
        "security_m6": {},
        "cleanroom_audit": {},
        "performance": {},
    }

    # --- A. Pure Logic Unit Tests ---
    sp = SpringSolver(0.0, response=0.45, damping_ratio=0.80)
    sp.set_target(1.0)
    peak = 0.0
    for _ in range(180):
        v = sp.step(1.0 / 120.0)
        peak = max(peak, v)
    assert sp.is_settled(), "Spring must settle within 1.5s"
    assert 1.005 <= peak <= 1.08, f"Expected slight water-like overshoot, got {peak}"
    results["unit_tests"]["spring_solver_overshoot_and_settle"] = f"PASS (peak={peak:.4f}, settled={sp.position:.4f})"

    # Ring placement on edges, corners, multi-monitor, and 100/125/150/200% DPI
    monitors = [
        (0, 0, 1920, 1040),
        (1920, 0, 3840, 1080),
    ]
    for mon in monitors:
        for dpi in (1.0, 1.25, 1.5, 2.0):
            r_px = round(230.0 * dpi)
            for cx_in, cy_in in [(mon[0] + 5, mon[1] + 5), (mon[2] - 5, mon[3] - 5), ((mon[0]+mon[2])//2, (mon[1]+mon[3])//2)]:
                rx, ry = compute_ring_placement(cx_in, cy_in, mon, dpi)
                assert mon[0] + r_px <= rx <= mon[2] - r_px
                assert mon[1] + r_px <= ry <= mon[3] - r_px
                assert math.hypot(rx - cx_in, ry - cy_in) >= r_px
    results["unit_tests"]["ring_placement_multi_monitor_dpi"] = "PASS (tested 2 monitors x 4 DPI scales x corners/centres)"

    # 360-degree clockwise and counter-clockwise rotation paging (multi-page AND single-page)
    pager = RotationPager(13, deadzone_radius=38.0)
    flipped_cw = None
    max_cw_prog = 0.0
    for step_deg in range(0, 390, 5):
        rad = math.radians(step_deg)
        p = pager.update_pointer(120.0 * math.sin(rad), -120.0 * math.cos(rad))
        max_cw_prog = max(max_cw_prog, pager.accumulated_rad / (2.0 * math.pi))
        if p is not None:
            flipped_cw = p
            break
    assert flipped_cw == 1 and max_cw_prog > 0.90, f"Expected page 1 after CW 360 turn, got {flipped_cw}"

    flipped_ccw = None
    min_ccw_prog = 0.0
    for step_deg in range(385, -30, -5):
        rad = math.radians(step_deg)
        p = pager.update_pointer(120.0 * math.sin(rad), -120.0 * math.cos(rad))
        min_ccw_prog = min(min_ccw_prog, pager.accumulated_rad / (2.0 * math.pi))
        if p is not None:
            flipped_ccw = p
            break
    assert flipped_ccw == 0 and min_ccw_prog < -0.85, f"Expected page 0 after CCW 360 turn, got {flipped_ccw}"

    pager_single = RotationPager(6, deadzone_radius=38.0)
    for step_deg in range(0, 180, 10):
        rad = math.radians(step_deg)
        pager_single.update_pointer(120.0 * math.sin(rad), -120.0 * math.cos(rad))
    assert pager_single.accumulated_rad > 2.5, "Single-page ring must also track 360° cursor rotation arc"
    results["unit_tests"]["rotation_paging_cw_ccw_hysteresis"] = "PASS (CW -> Page 1, CCW -> Page 0, single-page & multi-page 360° arc verified)"

    # Compatibility filtering & usage ranking tests (PROMPT.md §4 & §6)
    ranked_png = filter_and_rank_tools([Path("sample.png")], ffmpeg_available=True, usage_counts={"img_to_webp": 15, "Edit image": 20})
    convert_group = dict(ranked_png).get("Convert image", [])
    convert_ids = [t[0] for t in convert_group]
    assert "img_to_png" not in convert_ids, "Same-format PNG->PNG must be hidden when staging a .png file"
    assert convert_ids[0] == "img_to_webp", "Most-used tool img_to_webp must rank first on Page 1"
    ranked_no_ff = filter_and_rank_tools([Path("clip.mp4")], ffmpeg_available=False)
    assert "Video & audio" not in dict(ranked_no_ff), "Video & audio group must hide when FFmpeg is absent"
    assert parse_pages("1-3,5,7-", 9) == [1, 2, 3, 5, 7, 8, 9]
    assert abs(parse_time("01:02.5") - 62.5) < 1e-6
    results["unit_tests"]["tool_filtering_usage_ranking_and_parsers"] = "PASS (same-format exclusion, FFmpeg toggle, usage ordering, page & time parsers verified)"

    # --- A2. Crystal-Clear Liquid-Glass & Zero-Orange-Tint Optical Audit ---
    from liquid_optics import (
        render_optical_ring_rgba,
        render_optical_window_rgba,
        render_optical_island_rgba,
    )
    ring_cw = render_optical_ring_rgba(
        ["Convert image", "Edit image", "PDF & docs", "Video & audio", "Archives", "Data & text"],
        hover_seg=1, turn_progress=0.45, staged_count=3, subtitle="Edit image (Pg 1/2)", light_angle_deg=60.0
    )
    ring_ccw = render_optical_ring_rgba(
        ["Convert image", "Edit image", "PDF & docs", "Video & audio", "Archives", "Data & text"],
        hover_seg=4, turn_progress=-0.45, staged_count=3, subtitle="Edit image (Pg 2/2)", light_angle_deg=240.0
    )
    win_opt = render_optical_window_rgba(
        "Project_Files.zip", "2.4 GB", "Compress", "Quality (1-100)", mode="options", quality_val=85, chip_idx=1
    )
    win_prog = render_optical_window_rgba(
        "Project_Files.zip", "2.4 GB", "Compress", "Quality (1-100)", mode="progress", progress_frac=0.68
    )
    island_clear = render_optical_island_rgba(
        "Project_Files", "18 files  •  2.4 GB", progress_frac=0.68, paused=False
    )

    for surf_name, surf_im in [
        ("circle_ring_cw_360", ring_cw),
        ("circle_ring_ccw_360", ring_ccw),
        ("options_window_liquid_glass", win_opt),
        ("progress_window_liquid_glass", win_prog),
        ("dynamic_island_liquid_glass", island_clear),
    ]:
        arr = np.asarray(surf_im, dtype=np.float32)
        vis = arr[..., 3] > 20
        r_ch, g_ch, b_ch = arr[vis, 0], arr[vis, 1], arr[vis, 2]
        orange_pixels = int(np.sum((r_ch - b_ch > 18.0) & (r_ch > 140.0) & (b_ch < 120.0)))
        assert orange_pixels == 0, f"Found {orange_pixels} orange-tinted pixels in {surf_name}!"
        mean_rgb = (round(float(np.mean(r_ch)), 1), round(float(np.mean(g_ch)), 1), round(float(np.mean(b_ch)), 1))
        assert mean_rgb[2] >= mean_rgb[0], f"Expected cool crystal-clear glass (B >= R) in {surf_name}, got {mean_rgb}"
        results["liquid_glass_optics_and_zero_orange_audit"][surf_name] = (
            f"PASS (size={surf_im.size}, orange_pixels=0, mean_rgb={mean_rgb})"
        )

    # --- B. All 50 Tools Integration Suite ---
    with tempfile.TemporaryDirectory() as tmp:
        tdir = Path(tmp)

        img_path = tdir / "photo.png"
        im = Image.new("RGBA", (320, 240), (40, 120, 220, 255))
        d = ImageDraw.Draw(im)
        d.rectangle([40, 40, 280, 200], fill=(180, 225, 255, 255))
        im.save(img_path, "PNG")

        jpg_path = tdir / "camera.jpg"
        im.convert("RGB").save(jpg_path, "JPEG", quality=92)

        pdf1 = tdir / "doc1.pdf"
        c1 = rl_canvas.Canvas(str(pdf1), pagesize=A4)
        c1.setFont("Helvetica-Bold", 16)
        c1.drawString(72, 760, "PROJECT OVERVIEW")
        c1.setFont("Helvetica", 12)
        c1.drawString(72, 730, "First page paragraph text for extraction and conversion.")
        c1.drawImage(str(img_path), 72, 560, width=120, height=90)
        c1.showPage()
        c1.setFont("Helvetica", 12)
        c1.drawString(72, 760, "Second page content for split and reverse tests.")
        c1.save()

        pdf2 = tdir / "doc2.pdf"
        c2 = rl_canvas.Canvas(str(pdf2), pagesize=A4)
        c2.setFont("Helvetica", 12)
        c2.drawString(72, 760, "Appendix page from second PDF document.")
        c2.save()

        docx_path = tdir / "sample_doc.docx"
        wdoc = docx.Document()
        wdoc.add_heading("Clean Room Word Document", level=1)
        p_bold = wdoc.add_paragraph()
        r_bold = p_bold.add_run("This paragraph verifies native Word to PDF layout with embedded images.")
        r_bold.bold = True
        wdoc.add_picture(str(img_path), width=Inches(1.8))
        tbl = wdoc.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = "Metric"
        tbl.rows[0].cells[1].text = "Value"
        tbl.rows[1].cells[0].text = "RAM Idle"
        tbl.rows[1].cells[1].text = "0 MB"
        wdoc.save(str(docx_path))

        vid_path = tdir / "clip.mp4"
        subprocess.run([
            _find_ffmpeg(), "-nostdin", "-y", "-hide_banner", "-loglevel", "error",
            "-f", "lavfi", "-i", "testsrc=duration=1.5:size=160x120:rate=15",
            "-f", "lavfi", "-i", "sine=frequency=440:duration=1.5",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest",
            str(vid_path)
        ], check=True)

        json_path = tdir / "records.json"
        json_path.write_text('[{"name": "Alpha", "score": "98"}, {"name": "Beta", "score": "95", "extra": "ok"}]', encoding="utf-8")

        csv_path = tdir / "table.csv"
        csv_path.write_text("name;score\nAlpha;98\nBeta;95\n", encoding="utf-8")

        txt_path = tdir / "notes.txt"
        txt_path.write_bytes("Line 1: Café\r\nLine 2: Résumé\r\n".encode("cp1252"))

        srt_path = tdir / "subs.srt"
        srt_path.write_text("1\n00:00:01,000 --> 00:00:03,500\nHello Dewdrop!\n", encoding="utf-8")

        vtt_path = tdir / "subs.vtt"
        vtt_path.write_text("WEBVTT\n\n00:00:01.000 --> 00:00:03.500\nHello Dewdrop!\n", encoding="utf-8")

        tool_invocations = [
            ("img_to_png", [jpg_path], {"width": 200}),
            ("img_to_jpeg", [img_path], {"quality": 85}),
            ("img_to_webp", [img_path], {"quality": 85}),
            ("img_to_gif", [img_path], {}),
            ("img_to_ico", [img_path], {}),
            ("img_compress", [jpg_path], {"target_kb": 15}),
            ("img_resize", [jpg_path], {"resize_mode": "percent", "percent": 50}),
            ("img_crop", [jpg_path], {"ratio": "1:1"}),
            ("img_rotate_flip", [jpg_path], {"rotation": 90, "flip": "h"}),
            ("img_effects", [jpg_path], {"effect": "bw"}),
            ("img_watermark", [jpg_path], {"text": "VERIFIED", "position": "tl"}),
            ("img_strip_meta", [jpg_path], {}),
            ("pdf_from_images", [img_path, jpg_path], {"combine": True, "page": "A4"}),
            ("pdf_merge", [pdf1, pdf2], {}),
            ("pdf_split", [pdf1], {}),
            ("pdf_extract_pages", [pdf1], {"pages": "1-1"}),
            ("pdf_delete_pages", [pdf1], {"pages": "2-"}),
            ("pdf_reverse", [pdf1], {}),
            ("pdf_rotate", [pdf1], {"rotation": 90, "pages": "1"}),
            ("pdf_compress", [pdf1], {"level": 3}),
            ("pdf_to_text", [pdf1], {}),
            ("pdf_extract_images", [pdf1], {}),
            ("pdf_strip_meta", [pdf1], {}),
            ("pdf_to_word", [pdf1], {}),
            ("word_to_pdf", [docx_path], {}),
            ("media_to_mp4", [vid_path], {}),
            ("media_to_webm", [vid_path], {}),
            ("media_to_mp3", [vid_path], {}),
            ("media_to_m4a", [vid_path], {}),
            ("media_to_wav", [vid_path], {}),
            ("media_to_flac", [vid_path], {}),
            ("video_compress", [vid_path], {"level": 2, "maxh": 360}),
            ("audio_compress", [vid_path], {"bitrate": 96}),
            ("media_trim", [vid_path], {"start": "00:00", "end": "00:01"}),
            ("video_to_gif", [vid_path], {"width": 160, "duration": 1.0}),
            ("video_mute", [vid_path], {}),
            ("video_frame_png", [vid_path], {"start": "00:00.2"}),
            ("archive_zip", [img_path, txt_path], {"level": 2, "combine": True}),
            ("archive_targz", [img_path, txt_path], {"level": 2}),
            ("archive_gzip", [txt_path], {"level": 2}),
        ]

        for tid, in_files, opts in tool_invocations:
            out_p = execute_tool(tid, in_files, opts)
            assert out_p.exists(), f"Tool {tid} did not create {out_p}"
            if tid in ("img_resize", "img_crop", "img_rotate_flip", "img_effects", "img_watermark", "img_strip_meta"):
                assert out_p.suffix.lower() == ".jpg", f"Edit image tool {tid} must preserve input .jpg format, got {out_p.suffix}"
            results["tools_50"][tid] = f"PASS -> {out_p.parent.name}/{out_p.name} ({out_p.stat().st_size if out_p.is_file() else 'dir'} B)"

        zip_created = tdir / "Dew Drop" / "photo.zip"
        ext_out = execute_tool("archive_extract", [zip_created], {})
        assert ext_out.is_dir() and any(ext_out.iterdir())
        results["tools_50"]["archive_extract"] = f"PASS -> {ext_out.parent.name}/{ext_out.name} (extracted {len(list(ext_out.iterdir()))} files)"

        data_tools = [
            ("json_pretty", [json_path], {}),
            ("json_minify", [json_path], {}),
            ("json_to_csv", [json_path], {}),
            ("csv_to_json", [csv_path], {}),
            ("text_utf8", [txt_path], {}),
            ("text_crlf", [txt_path], {}),
            ("srt_to_vtt", [srt_path], {}),
            ("vtt_to_srt", [vtt_path], {}),
            ("checksums_all", [img_path, txt_path], {}),
        ]
        for tid, in_files, opts in data_tools:
            out_p = execute_tool(tid, in_files, opts)
            assert out_p.exists() and out_p.stat().st_size > 0
            results["tools_50"][tid] = f"PASS -> {out_p.parent.name}/{out_p.name} ({out_p.stat().st_size} B)"

        assert len(results["tools_50"]) == 50, f"Expected 50 tools, tested {len(results['tools_50'])}"

        # Also test subprocess worker_main protocol
        job_json = tdir / "job.json"
        res_json = tdir / "res.json"
        job_json.write_text(json.dumps({
            "tool_id": "img_to_webp",
            "files": [str(img_path)],
            "options": {"quality": 80},
        }), encoding="utf-8")
        rc = worker_main(str(job_json), str(res_json))
        res_data = json.loads(res_json.read_text(encoding="utf-8"))
        assert rc == 0 and res_data["finished"] and len(res_data["outputs"]) == 1
        results["unit_tests"]["worker_subprocess_protocol"] = "PASS (worker_main atomic progress & completion JSON verified)"

        # --- C. 10-Sample PDF-to-Word Acceptance Test ---
        sample_specs = [
            ("01_plain_paragraphs", ["Plain paragraph one.", "Plain paragraph two with standard text."]),
            ("02_headings_hierarchy", ["MAIN CHAPTER HEADING", "Section body text under main heading.", "SUBSECTION HEADING", "More details."]),
            ("03_bullet_lists", ["SHOPPING AND TASK LIST", "- First bullet item", "- Second bullet item", "• Third bullet item"]),
            ("04_simple_table", ["QUARTERLY METRICS", "Quarter | Revenue | Margin", "Q1 | 1200 | 32%", "Q2 | 1450 | 35%"]),
            ("05_two_column_layout", ["LEFT AND RIGHT COLUMNS", "Column A line 1    Column B line 1", "Column A line 2    Column B line 2"]),
            ("06_embedded_image_caption", ["FIGURE REPORT", "Below is the embedded test diagram.", "Figure 1: Water droplet diagram."]),
            ("07_multipage_report", ["PAGE ONE SUMMARY", "Executive overview text.", "PAGE TWO DETAILS", "Detailed breakdown on page two."]),
            ("08_mixed_typography", ["TECHNICAL SPECIFICATION", "Parameter: Response=0.45s, Damping=0.80.", "- Verified spring physics"]),
            ("09_invoice_schedule", ["INVOICE STATEMENT", "Item | Qty | Price", "License | 1 | 0.00"]),
            ("10_release_checklist", ["FINAL GATE CHECKLIST", "- Fresh install verified", "- Zero background processes when idle"]),
        ]
        for s_name, lines in sample_specs:
            spdf = tdir / f"{s_name}.pdf"
            sc = rl_canvas.Canvas(str(spdf), pagesize=letter)
            y = 740
            for ln in lines:
                sc.setFont("Helvetica-Bold" if ln.isupper() else "Helvetica", 13 if ln.isupper() else 11)
                sc.drawString(64, y, ln)
                y -= 22
            if "image" in s_name:
                sc.drawImage(str(img_path), 64, y - 100, width=120, height=90)
            sc.save()
            out_docx = execute_tool("pdf_to_word", [spdf], {})
            parsed = docx.Document(str(out_docx))
            para_count = len(parsed.paragraphs) + len(parsed.tables)
            assert para_count >= 2, f"Sample {s_name} produced empty docx"
            results["pdf_to_word_10_samples"][s_name] = f"PASS ({para_count} structural blocks verified in Word .docx)"

        # --- D. 100-File Batch Test with Deliberate Failures ---
        batch_dir = tdir / "batch100"
        batch_dir.mkdir()
        pass_count = 0
        fail_count = 0
        t0 = time.perf_counter()
        for i in range(100):
            bf = batch_dir / f"item_{i:03d}.png"
            if i in (17, 42, 88):
                bf.write_bytes(b"NOT_A_VALID_PNG_HEADER_CORRUPT")
            else:
                im_small = Image.new("RGB", (64, 64), (i * 2 % 255, 100, 180))
                im_small.save(bf, "PNG")
            try:
                execute_tool("img_to_jpeg", [bf], {"quality": 80})
                pass_count += 1
            except Exception:
                fail_count += 1
        batch_elapsed = time.perf_counter() - t0
        assert pass_count == 97 and fail_count == 3
        results["batch_100"] = {
            "total": 100,
            "succeeded": pass_count,
            "failed_cleanly": fail_count,
            "elapsed_sec": round(batch_elapsed, 3),
            "status": "PASS",
        }

        # --- E. M6 Malicious Sample Suite ---
        zslip = tdir / "malicious_zipslip.zip"
        with zipfile.ZipFile(zslip, "w") as zf:
            zf.writestr("../../outside_escape.txt", "malicious payload")
        try:
            execute_tool("archive_extract", [zslip], {})
            raise AssertionError("Zip-slip was not refused!")
        except ValueError as e:
            results["security_m6"]["1_zip_slip_traversal"] = f"PASS (Refused: {e})"

        zbomb = tdir / "malicious_bomb.zip"
        with zipfile.ZipFile(zbomb, "w") as zf:
            for idx in range(2005):
                zf.writestr(f"f_{idx}.txt", "x")
        try:
            execute_tool("archive_extract", [zbomb], {})
            raise AssertionError("Zip bomb was not refused!")
        except ValueError as e:
            results["security_m6"]["2_zip_decompression_bomb"] = f"PASS (Refused: {e})"

        giant_png = tdir / "giant_100mp.png"
        Image.MAX_IMAGE_PIXELS = None
        Image.new("L", (10000, 10000), 0).save(giant_png, "PNG")
        Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
        try:
            execute_tool("img_to_jpeg", [giant_png], {})
            raise AssertionError("100 MP image bomb was not refused!")
        except Exception as e:
            results["security_m6"]["3_giant_40mp_image_bomb"] = f"PASS (Refused: {e})"

        js_pdf = tdir / "malicious_js.pdf"
        jw = pypdf.PdfWriter()
        jw.add_blank_page(width=200, height=200)
        jw.add_js("app.alert('XSS');")
        with open(js_pdf, "wb") as f:
            jw.write(f)
        cleaned_pdf = execute_tool("pdf_strip_meta", [js_pdf], {})
        raw_clean = cleaned_pdf.read_bytes()
        assert b"app.alert" not in raw_clean and b"/JavaScript" not in raw_clean and b"/OpenAction" not in raw_clean
        results["security_m6"]["4_pdf_javascript_autoaction_strip"] = "PASS (Verified /JavaScript and /OpenAction completely stripped)"

        bad_m3u8 = tdir / "crafted.m3u8"
        bad_m3u8.write_text("#EXTM3U\n#EXTINF:1.0,\nhttp://127.0.0.1:9999/secret.ts\n", encoding="utf-8")
        try:
            execute_tool("media_to_mp4", [bad_m3u8], {})
            raise AssertionError("Crafted network playlist was not refused!")
        except Exception as e:
            results["security_m6"]["5_ffmpeg_protocol_whitelist_lockdown"] = f"PASS (Blocked external protocol: {str(e)[:120]})"

        assert sanitize_stem("-rf --no-preserve-root") == "rf --no-preserve-root"
        assert sanitize_stem("CON.txt").startswith("_CON")
        assert sanitize_stem("NUL...") == "_NUL"
        results["security_m6"]["6_odd_filenames_and_device_names"] = "PASS (Leading '-', CON/NUL, and trailing dots neutralized)"

        macro_docx = tdir / "macro_trojan.docx"
        shutil.copyfile(docx_path, macro_docx)
        with zipfile.ZipFile(macro_docx, "a") as zf:
            zf.writestr("word/vbaProject.bin", b"FAKE_VBA_MACRO_STREAM")
        try:
            execute_tool("word_to_pdf", [macro_docx], {})
            raise AssertionError("Macro docx was not refused!")
        except ValueError as e:
            results["security_m6"]["7_macro_enabled_docx_refused"] = f"PASS (Refused: {e})"

    # --- F. Clean-Room & Zero-Legacy Identity Audit ---
    forbidden_patterns = [
        "Open" + "Tangerine",
        "Open " + "Tangerine",
        "Sat" + "suma",
        "Sintax" + " Error",
        "ply" + "ght",
        "SPDX-" + "License-Identifier",
        "GNU GENERAL " + "PUBLIC LICENSE",
        "A2A02332-" + "784A-4C0C-99DA-C82D08D2BDC8",
    ]
    scanned_files = 0
    hits = []
    for p in ROOT.rglob("*"):
        if p.is_file() and p.suffix in (".rs", ".toml", ".iss", ".ps1", ".py", ".md", ".manifest"):
            scanned_files += 1
            txt = p.read_text(encoding="utf-8", errors="ignore")
            for pat in forbidden_patterns:
                if pat.lower() in txt.lower():
                    hits.append(f"{p.relative_to(ROOT)}: {pat}")
    assert not hits, f"Clean-room audit failed with hits: {hits}"
    results["cleanroom_audit"] = {
        "files_scanned": scanned_files,
        "forbidden_matches": len(hits),
        "license_files_present": (ROOT / "LICENSE").exists() or (ROOT / "NOTICE").exists(),
        "status": "PASS (0 legacy strings, 0 automatic LICENSE/NOTICE files)",
    }

    out_json = EVIDENCE_DIR / "test_suite_results.json"
    out_json.write_text(json.dumps(results, indent=2), encoding="utf-8")
    print(f"All verification suites PASSED: 50/50 tools, 10/10 PDF->Word samples, 100-file batch, 7/7 M6 security guards, Crystal-Clear Liquid-Glass zero-orange audit, Clean-Room audit ({scanned_files} files).")


if __name__ == "__main__":
    run_all_tests()
